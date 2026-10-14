"""
Reusable Reporting & Documentation engine — the asset owned by the 'Reporting & Documentation Agent'.
Give it a structured report spec; it produces a nicely-styled, comprehensive document in md / docx / pdf /
pptx, with brand (logo + company info). Reusable by ANY project (SOC health report, meeting report, deploy
report, …). Outputs land under the project workspace so they flow through the SAME chat pool.

spec = {
  "title": str, "subtitle": str,
  "brand": {"name": str, "info": str, "logo": path, "primary": "#2B5CE0", "accent": "#1F9A61"},
  "meta": [("Label","value"), ...],
  "sections": [
     {"heading": str, "level": 1|2, "text": str, "bullets": [..], "badge": ("High","#c33"),
      "table": {"headers": [...], "rows": [[...], ...]}}
  ],
  "footer": str,
}
"""
from __future__ import annotations
import io, os, re, datetime, pathlib

DEFAULT_BRAND = {
    "name": "Odin AI Systems", "info": "Odin-Orchestrator · Autonomous Multi-Agent Platform",
    "logo": str(pathlib.Path(__file__).resolve().parent.parent / "desktop" / "build" / "icon.png"),
    "primary": "#2B5CE0", "accent": "#1F9A61", "dark": "#222738",
}

def _brand(spec): b = dict(DEFAULT_BRAND); b.update(spec.get("brand") or {}); return b

# ------------------------------------------------------------------ Markdown
def to_md(spec) -> str:
    b = _brand(spec)
    out = [f"# {spec.get('title','Report')}", ""]
    if spec.get("subtitle"): out += [f"_{spec['subtitle']}_", ""]
    out += [f"**{b['name']}** — {b['info']}", ""]
    for k, v in (spec.get("meta") or []): out.append(f"- **{k}:** {v}")
    out.append("")
    for s in spec.get("sections", []):
        h = "#" * (1 + int(s.get("level", 1)))
        badge = f"  `{s['badge'][0]}`" if s.get("badge") else ""
        out += [f"{h} {s.get('heading','')}{badge}", ""]
        if s.get("text"): out += [s["text"], ""]
        for x in s.get("bullets", []): out.append(f"- {x}")
        if s.get("bullets"): out.append("")
        t = s.get("table")
        if t:
            out.append("| " + " | ".join(t["headers"]) + " |")
            out.append("| " + " | ".join("---" for _ in t["headers"]) + " |")
            for row in t["rows"]:
                out.append("| " + " | ".join(str(c) for c in row) + " |")
            out.append("")
    if spec.get("footer"): out += ["---", spec["footer"]]
    return "\n".join(out) + "\n"

def _hex(c): c = c.lstrip("#"); return tuple(int(c[i:i+2], 16) for i in (0, 2, 4))

# ------------------------------------------------------------------ DOCX
def to_docx(spec) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    b = _brand(spec); pr = RGBColor(*_hex(b["primary"])); ac = RGBColor(*_hex(b["accent"]))
    def shade(p, fill):
        e = OxmlElement("w:shd"); e.set(qn("w:val"), "clear"); e.set(qn("w:fill"), fill)
        p._p.get_or_add_pPr().append(e)
    doc = Document(); doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)
    # header band with logo
    if b.get("logo") and os.path.exists(b["logo"]):
        try: doc.add_picture(b["logo"], width=Inches(0.55))
        except Exception: pass
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(spec.get("title", "Report")); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = pr
    shade(t, "EAF0FF")
    if spec.get("subtitle"):
        st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = st.add_run(spec["subtitle"]); sr.italic = True; sr.font.size = Pt(9.5)
    bp = doc.add_paragraph(); bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bp.add_run(f"{b['name']} — {b['info']}"); br.font.size = Pt(8.5); br.font.color.rgb = RGBColor(*_hex(b["dark"]))
    for k, v in (spec.get("meta") or []):
        m = doc.add_paragraph(); m.add_run(f"{k}: ").bold = True; m.add_run(str(v))
    doc.add_paragraph("")
    for s in spec.get("sections", []):
        h = doc.add_paragraph(); hr = h.add_run(s.get("heading", "")); hr.bold = True
        hr.font.size = Pt(14 if s.get("level", 1) <= 1 else 12); hr.font.color.rgb = pr if s.get("level", 1) <= 1 else ac
        if s.get("level", 1) <= 1: shade(h, "F2F6FF")
        if s.get("badge"):
            bd = doc.add_paragraph(); bdr = bd.add_run(f" {s['badge'][0]} "); bdr.bold = True
            bdr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(bd, s['badge'][1].lstrip('#'))
        if s.get("text"): doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", s["text"]))
        for x in s.get("bullets", []): doc.add_paragraph(str(x), style="List Bullet")
        tb = s.get("table")
        if tb:
            table = doc.add_table(rows=1, cols=len(tb["headers"])); table.style = "Light Grid Accent 1"
            for i, hcell in enumerate(tb["headers"]):
                cp = table.rows[0].cells[i].paragraphs[0]; cr = cp.add_run(str(hcell)); cr.bold = True
            for row in tb["rows"]:
                cells = table.add_row().cells
                for i, c in enumerate(row): cells[i].text = str(c)
    if spec.get("footer"):
        f = doc.add_paragraph(); fr = f.add_run(spec["footer"]); fr.font.size = Pt(8); fr.italic = True
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

# ------------------------------------------------------------------ PDF
def to_pdf(spec) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image,
                                    ListFlowable, ListItem)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    b = _brand(spec); primary = colors.HexColor(b["primary"]); accent = colors.HexColor(b["accent"])
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], textColor=primary, fontSize=15, spaceBefore=10)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=accent, fontSize=12)
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=ss["BodyText"], fontSize=8, textColor=colors.HexColor(b["dark"]))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16*mm, bottomMargin=14*mm, leftMargin=16*mm, rightMargin=16*mm,
                            title=spec.get("title", "Report"), author=b["name"])
    story = []
    if b.get("logo") and os.path.exists(b["logo"]):
        try: story.append(Image(b["logo"], width=14*mm, height=14*mm))
        except Exception: pass
    story.append(Paragraph(f"<b>{spec.get('title','Report')}</b>", ParagraphStyle("title", parent=h1, fontSize=22, textColor=primary)))
    if spec.get("subtitle"): story.append(Paragraph(f"<i>{spec['subtitle']}</i>", small))
    story.append(Paragraph(f"{b['name']} — {b['info']}", small))
    if spec.get("meta"):
        story.append(Spacer(1, 4))
        story.append(Table([[Paragraph(f"<b>{k}</b>", small), Paragraph(str(v), small)] for k, v in spec["meta"]],
                           colWidths=[45*mm, 120*mm],
                           style=TableStyle([("BOX", (0,0), (-1,-1), 0.4, colors.HexColor("#DDE3F0")),
                                             ("INNERGRID", (0,0), (-1,-1), 0.3, colors.HexColor("#EEF2FB"))])))
    story.append(Spacer(1, 8))
    for s in spec.get("sections", []):
        head = s.get("heading", "")
        if s.get("badge"): head += f' &nbsp;<font backColor="{s["badge"][1]}" color="white"> {s["badge"][0]} </font>'
        story.append(Paragraph(head, h1 if s.get("level", 1) <= 1 else h2))
        if s.get("text"): story.append(Paragraph(re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s["text"]), body))
        if s.get("bullets"):
            story.append(ListFlowable([ListItem(Paragraph(str(x), body)) for x in s["bullets"]], bulletType="bullet"))
        tb = s.get("table")
        if tb:
            data = [[Paragraph(f"<b>{h}</b>", small) for h in tb["headers"]]] + \
                   [[Paragraph(str(c), small) for c in row] for row in tb["rows"]]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), primary), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                                   ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#DDE3F0")),
                                   ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F6F8FE")])]))
            story.append(t)
        story.append(Spacer(1, 6))
    if spec.get("footer"): story.append(Paragraph(spec["footer"], small))
    doc.build(story); return buf.getvalue()

# ------------------------------------------------------------------ PPTX
def to_pptx(spec) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    b = _brand(spec); pr = RGBColor(*_hex(b["primary"]))
    prs = Presentation(); blank = prs.slide_layouts[6]
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = spec.get("title", "Report")
    try: s0.placeholders[1].text = f"{spec.get('subtitle','')}\n{b['name']} — {b['info']}"
    except Exception: pass
    if b.get("logo") and os.path.exists(b["logo"]):
        try: s0.shapes.add_picture(b["logo"], Inches(8.4), Inches(0.3), height=Inches(0.9))
        except Exception: pass
    for s in spec.get("sections", []):
        sl = prs.slides.add_slide(blank)
        tb = sl.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(0.9)).text_frame
        tb.text = s.get("heading", ""); tb.paragraphs[0].font.size = Pt(28); tb.paragraphs[0].font.bold = True
        tb.paragraphs[0].font.color.rgb = pr
        body = sl.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(9), Inches(5)).text_frame; body.word_wrap = True
        if s.get("text"): body.text = re.sub(r"\*\*(.+?)\*\*", r"\1", s["text"])
        for x in s.get("bullets", []):
            p = body.add_paragraph(); p.text = f"• {x}"; p.font.size = Pt(14)
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()

EXPORTERS = {"md": (to_md, "text/markdown", "md"), "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
             "pdf": (to_pdf, "application/pdf", "pdf"),
             "pptx": (to_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx")}

def render(spec: dict, fmt: str):
    fmt = fmt if fmt in EXPORTERS else "md"
    fn, mime, ext = EXPORTERS[fmt]
    data = fn(spec)
    if isinstance(data, str): data = data.encode("utf-8")
    return data, mime, ext

def render_to_file(spec: dict, fmt: str, out_dir: pathlib.Path, stem: str) -> pathlib.Path:
    data, mime, ext = render(spec, fmt)
    out_dir.mkdir(parents=True, exist_ok=True)
    p = out_dir / f"{stem}.{ext}"; p.write_bytes(data); return p
